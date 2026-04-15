# indexer.py — Crawl directories and store file metadata + content in ChromaDB.
"""
Every file encountered gets a searchable metadata record.
DOCX / XLSX / XLS / PDF files also get text-chunk records for semantic search.

Change detection uses os.stat().st_mtime_ns (integer nanoseconds) — fast, no
file reads. Only files whose modification timestamp changed are re-indexed.

Usage
-----
  python indexer.py              # index all WATCH_PATHS from config.py
  python indexer.py --force      # re-index everything ignoring cached mtimes
  python indexer.py --cleanup    # remove records for files no longer on disk
  python indexer.py --stats      # print collection statistics
  python indexer.py --file PATH  # index (or re-index) a single file
"""

from __future__ import annotations

import argparse
import hashlib
import logging
import os
import sys
from datetime import datetime
from pathlib import Path

import chromadb
from chromadb.utils.embedding_functions import SentenceTransformerEmbeddingFunction

from config import (
    WATCH_PATHS,
    SKIP_DIRS,
    SKIP_FILES,
    CONTENT_EXTENSIONS,
    CHUNK_SIZE,
    CHUNK_OVERLAP,
    EMBED_MODEL,
    DB_PATH,
    COLLECTION_NAME,
    MAX_FILE_SIZE_MB,
)

# ── Logging ───────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler(
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "indexer.log"),
            encoding="utf-8",
        ),
    ],
)
log = logging.getLogger(__name__)

# ── ChromaDB setup ────────────────────────────────────────────────────────────
# SentenceTransformerEmbeddingFunction auto-embeds documents on add() and
# query_texts on query() — no manual embed() calls needed anywhere.
log.info(f"Loading embedding model '{EMBED_MODEL}'…")
_embed_fn = SentenceTransformerEmbeddingFunction(model_name=EMBED_MODEL)
_client = chromadb.PersistentClient(path=DB_PATH)
collection = _client.get_or_create_collection(
    name=COLLECTION_NAME,
    embedding_function=_embed_fn,
    metadata={"hnsw:space": "cosine"},  # cosine distance: 0=identical, 1=unrelated
)
log.info(f"Collection '{COLLECTION_NAME}' ready. {collection.count()} records.")


# ── Stable file ID ────────────────────────────────────────────────────────────

def _stable_id(path: str) -> str:
    """MD5 of the normalised absolute path — case-insensitive on Windows."""
    normed = os.path.normcase(os.path.abspath(path))
    return hashlib.md5(normed.encode()).hexdigest()


def _meta_id(path: str) -> str:
    return _stable_id(path) + "_meta"


def _chunk_id(path: str, idx: int) -> str:
    return f"{_stable_id(path)}_chunk_{idx}"


# ── Change detection ──────────────────────────────────────────────────────────

def _needs_reindex(path: str) -> bool:
    """True if the file is new or its mtime_ns differs from the stored value."""
    mid = _meta_id(path)
    try:
        result = collection.get(ids=[mid], include=["metadatas"])
        if not result["ids"]:
            return True
        stored = result["metadatas"][0].get("mtime_ns")
        current = os.stat(path).st_mtime_ns
        return stored != current
    except Exception:
        return True


# ── Chunk helpers ─────────────────────────────────────────────────────────────

def _chunk_text(text: str) -> list[str]:
    """Split text into overlapping windows, preferring to break on whitespace."""
    text = text.strip()
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            for sep in ("\n", " "):
                pos = text.rfind(sep, start + CHUNK_SIZE // 2, end)
                if pos > start:
                    end = pos
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - CHUNK_OVERLAP
    return chunks


# ── Content parsers ───────────────────────────────────────────────────────────

def _parse_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"[Sheet: {sheet_name}]")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def _parse_xls(path: str) -> str:
    import xlrd
    wb = xlrd.open_workbook(path)
    parts: list[str] = []
    for sheet in wb.sheets():
        parts.append(f"[Sheet: {sheet.name}]")
        for rx in range(sheet.nrows):
            cells = [
                str(sheet.cell_value(rx, cx))
                for cx in range(sheet.ncols)
                if str(sheet.cell_value(rx, cx)).strip()
            ]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _parse_pdf(path: str) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(path)
    pages: list[str] = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            pages.append(f"[Page {i}]\n{text}")
    doc.close()
    return "\n\n".join(pages)


_PARSERS: dict[str, callable] = {
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".xls":  _parse_xls,
    ".pdf":  _parse_pdf,
}


# ── Metadata helpers ──────────────────────────────────────────────────────────

def _file_stat(path: str) -> dict:
    """OS-level file attributes. Uses st_ctime as creation time on Windows/NTFS."""
    stat = os.stat(path)
    p = Path(path).resolve()
    return {
        "file_name":    p.name,
        "file_path":    str(p),
        "extension":    p.suffix.lower(),
        "size_bytes":   stat.st_size,
        "date_created": datetime.fromtimestamp(stat.st_ctime).isoformat(timespec="seconds"),
        "date_modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "mtime_ns":     stat.st_mtime_ns,   # int — used for change detection
    }


def _meta_document(stat: dict, content_indexed: bool) -> str:
    """Human-readable text that gets embedded as the metadata record."""
    return (
        f"File: {stat['file_name']}\n"
        f"Path: {stat['file_path']}\n"
        f"Type: {stat['extension'] or '(no extension)'}\n"
        f"Size: {stat['size_bytes']:,} bytes\n"
        f"Created:  {stat['date_created']}\n"
        f"Modified: {stat['date_modified']}\n"
        f"Content indexed: {content_indexed}"
    )


# ── Delete helpers ────────────────────────────────────────────────────────────

def _delete_chunks(path: str, total_chunks: int) -> None:
    """Delete chunk records by reconstructed IDs — no where-filter scan needed."""
    if total_chunks <= 0:
        return
    fid = _stable_id(path)
    ids = [_chunk_id(path, i) for i in range(total_chunks)]
    # get() silently omits absent IDs; delete() raises on absent IDs → check first
    present = collection.get(ids=ids, include=[])
    if present["ids"]:
        collection.delete(ids=present["ids"])


def remove_file(path: str) -> None:
    """Remove all index records for *path*. Safe to call on un-indexed paths."""
    mid = _meta_id(path)
    try:
        result = collection.get(ids=[mid], include=["metadatas"])
        if not result["ids"]:
            return
        total_chunks = int(result["metadatas"][0].get("total_chunks", 0))
        _delete_chunks(path, total_chunks)
        collection.delete(ids=[mid])
        log.info(f"[removed] {result['metadatas'][0].get('file_name', path)}")
    except Exception as e:
        log.debug(f"remove_file({path}): {e}")


# ── Core indexing ─────────────────────────────────────────────────────────────

def index_file(path: str, force: bool = False) -> str:
    """
    Index a single file. Returns a status string:
      'skipped'   — mtime unchanged, no action taken
      'meta_only' — file catalogued but content not extracted
      'indexed'   — content extracted and stored as searchable chunks
      'error'     — stat or unexpected failure
    """
    try:
        path = str(Path(path).resolve())
    except Exception:
        pass

    try:
        stat = _file_stat(path)
    except OSError as exc:
        log.warning(f"Cannot stat '{path}': {exc}")
        return "error"

    if not force and not _needs_reindex(path):
        return "skipped"

    ext = stat["extension"]
    size_mb = stat["size_bytes"] / (1024 * 1024)
    is_content = ext in CONTENT_EXTENSIONS and size_mb <= MAX_FILE_SIZE_MB

    if not is_content:
        return _store_meta_only(path, stat)

    return _store_content(path, stat, ext)


def _store_meta_only(path: str, stat: dict) -> str:
    mid = _meta_id(path)
    # Remove any previous records (may have had content on a prior run)
    old = collection.get(ids=[mid], include=["metadatas"])
    if old["ids"]:
        old_chunks = int(old["metadatas"][0].get("total_chunks", 0))
        _delete_chunks(path, old_chunks)
        collection.delete(ids=[mid])

    collection.add(
        ids=[mid],
        documents=[_meta_document(stat, content_indexed=False)],
        metadatas=[{**stat, "record_type": "meta", "content_indexed": False,
                    "total_chunks": 0, "chunk_index": -1}],
    )
    return "meta_only"


def _store_content(path: str, stat: dict, ext: str) -> str:
    parser = _PARSERS[ext]
    try:
        text = parser(path)
    except Exception as exc:
        log.warning(f"Parse failed '{path}': {exc} — storing metadata only")
        return _store_meta_only(path, stat)

    chunks = _chunk_text(text)
    if not chunks:
        log.info(f"No text extracted from '{path}' — storing metadata only")
        return _store_meta_only(path, stat)

    mid = _meta_id(path)
    total = len(chunks)

    # Remove old records before writing new ones
    old = collection.get(ids=[mid], include=["metadatas"])
    if old["ids"]:
        old_chunks = int(old["metadatas"][0].get("total_chunks", 0))
        _delete_chunks(path, old_chunks)
        collection.delete(ids=[mid])

    # Write chunk records (in batches to avoid large single requests)
    BATCH = 100
    for b in range(0, total, BATCH):
        batch = chunks[b:b + BATCH]
        collection.add(
            ids=[_chunk_id(path, b + i) for i in range(len(batch))],
            documents=batch,
            metadatas=[{
                **stat,
                "record_type":     "chunk",
                "content_indexed": True,
                "total_chunks":    total,
                "chunk_index":     b + i,
            } for i in range(len(batch))],
        )

    # Write metadata record last (total_chunks needed for future deletions)
    collection.add(
        ids=[mid],
        documents=[_meta_document(stat, content_indexed=True)],
        metadatas=[{**stat, "record_type": "meta", "content_indexed": True,
                    "total_chunks": total, "chunk_index": -1}],
    )
    return "indexed"


# ── Full scan ─────────────────────────────────────────────────────────────────

def index_paths(paths: list[str], force: bool = False) -> dict[str, int]:
    counts: dict[str, int] = {
        "indexed": 0, "meta_only": 0, "skipped": 0, "error": 0
    }
    total = 0

    for base in paths:
        if not os.path.isdir(base):
            log.warning(f"Watch path not found, skipping: {base}")
            continue
        for root, dirs, files in os.walk(base, topdown=True):
            dirs[:] = [
                d for d in dirs
                if d not in SKIP_DIRS and not d.startswith(".")
            ]
            for fname in files:
                if fname in SKIP_FILES or fname.startswith("~$"):
                    continue
                result = index_file(os.path.join(root, fname), force=force)
                counts[result] += 1
                total += 1
                if total % 100 == 0:
                    log.info(f"Progress: {total} files — {counts}")

    log.info(f"Scan complete. {total} files — {counts}")
    return counts


# ── Cleanup ───────────────────────────────────────────────────────────────────

def cleanup_deleted() -> int:
    """Remove records for files that no longer exist on disk."""
    log.info("Scanning index for stale records…")
    removed = 0
    offset = 0
    PAGE = 500

    while True:
        try:
            page = collection.get(limit=PAGE, offset=offset, include=["metadatas"])
        except Exception as exc:
            log.warning(f"Cleanup page error: {exc}")
            break

        if not page["ids"]:
            break

        for record_id, meta in zip(page["ids"], page["metadatas"]):
            if meta.get("record_type") != "meta":
                offset += 1
                continue
            fpath = meta.get("file_path", "")
            if fpath and not os.path.exists(fpath):
                log.info(f"Removing stale record: {fpath}")
                remove_file(fpath)
                removed += 1
            else:
                offset += 1

    log.info(f"Cleanup done — {removed} stale file(s) removed.")
    return removed


# ── Stats ─────────────────────────────────────────────────────────────────────

def show_stats() -> None:
    total = collection.count()
    print(f"\nCollection '{COLLECTION_NAME}' at {DB_PATH}")
    print(f"{'─'*50}")
    print(f"  Total records (chunks + metadata): {total:,}")
    if total == 0:
        print("  (Index is empty — run 'python indexer.py' to populate.)")
        return

    all_records = collection.get(include=["metadatas"])
    meta_records  = [m for m in all_records["metadatas"] if m.get("record_type") == "meta"]
    chunk_records = [m for m in all_records["metadatas"] if m.get("record_type") == "chunk"]
    indexed   = [m for m in meta_records if m.get("content_indexed") is True]
    meta_only = [m for m in meta_records if not m.get("content_indexed")]

    from collections import Counter
    ext_counts = Counter(m.get("extension", "(none)") for m in meta_records)

    print(f"  Files tracked     : {len(meta_records):,}")
    print(f"  Content indexed   : {len(indexed):,}")
    print(f"  Metadata only     : {len(meta_only):,}")
    print(f"  Searchable chunks : {len(chunk_records):,}")
    print(f"\n  File types:")
    for ext, cnt in ext_counts.most_common(20):
        marker = " [content]" if ext in CONTENT_EXTENSIONS else ""
        print(f"    {(ext or '(none)'):<22} {cnt:>6,}{marker}")
    print()


# ── CLI entry point ───────────────────────────────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="FileIndexer — index local files into a ChromaDB vector database"
    )
    parser.add_argument(
        "--force", action="store_true",
        help="Re-index all files regardless of modification time"
    )
    parser.add_argument(
        "--cleanup", action="store_true",
        help="Remove index records for files that no longer exist on disk"
    )
    parser.add_argument(
        "--stats", action="store_true",
        help="Print collection statistics and exit"
    )
    parser.add_argument(
        "--file", metavar="PATH",
        help="Index (or re-index) a single specific file"
    )
    parser.add_argument(
        "--paths", nargs="+", metavar="PATH",
        help="Override WATCH_PATHS for this run"
    )
    args = parser.parse_args()

    if args.stats:
        show_stats()
    elif args.cleanup:
        cleanup_deleted()
    elif args.file:
        result = index_file(args.file, force=True)
        print(f"Result: {result}")
    else:
        paths = args.paths or WATCH_PATHS
        log.info(f"Indexing {len(paths)} path(s): {paths}")
        counts = index_paths(paths, force=args.force)
        print(f"\nDone: {counts}")
